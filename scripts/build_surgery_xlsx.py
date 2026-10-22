# -*- coding: utf-8 -*-
"""
수술비 3종(1-5종Ⅱ · 1-7종 · 130대/131대질병) — 보장 수술 · 질병 · KCD 정리 엑셀

원천 : tool.html (특약검색기)
  · 1-5종Ⅱ  : 수술비 비교 ④ 표(SC.detail) — 약관 「1-5종 수술분류표Ⅱ」와 대조 끝(종 99줄 일치)
  · 1-7종    : 약관 【별표3】 1-7종 수술분류표 정리본(DATA tables) — 약관 PDF 4권과 수술코드 630개 · 종 일치
  · 131대    : 수술비 비교 ⑥ 표(SC.d131) + 특약 마스터(케219 · 통123 세부보장 k/x) — 약관 분류표와 대조 끝
실행 : python3 scripts/build_surgery_xlsx.py [출력.xlsx]   (기본 dist/surgery_benefit_kcd_2610.xlsx)
"""
import io, json, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tooldata import PAT, inflate, ROOT
from same_riders import _expand, covers
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'dist', 'surgery_benefit_kcd_2610.xlsx')
FONT = 'Arial'
HEAD = PatternFill('solid', fgColor='0B7C6F')
SEC = PatternFill('solid', fgColor='E6F6F2')
THIN = Side(style='thin', color='D5DCE4')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load():
    html = io.open(os.path.join(ROOT, 'tool.html'), encoding='utf-8').read()
    D = inflate(json.loads(re.search(PAT, html, re.S).group(2)))
    blk = html[html.index('/* === 수술비 비교'):]
    SC, _ = json.JSONDecoder().raw_decode(blk[blk.index('var SC = ') + 9:])
    d17 = json.loads(re.search(r'SC\.d17=(\[.*?\]);\n', blk).group(1))
    d131 = json.loads(re.search(r'SC\.d131=(\[.*?\]);\n', blk).group(1))
    return D, SC, d17, d131


def sheet(wb, title, headers, rows, widths, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=FONT, size=10, italic=True, color='46536B')
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.row_dimensions[1].height = 30
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical='center')
        r0 = 2
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font = Font(name=FONT, bold=True, color='FFFFFF'); c.fill = HEAD; c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=FONT, size=10); c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=True)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = '%s:%s%d' % (ws.cell(r0, 1).coordinate, get_column_letter(len(headers)), r0 + len(rows))
    return ws


def main():
    D, SC, d17, d131 = load()
    by = {r['id']: r for r in D['riders']}
    cn = D.get('codenames') or {}
    wb = Workbook(); wb.remove(wb.active)

    # ── 1-5종Ⅱ ──
    rows15 = [[r[0], r[1], r[2], r[3], r[4]] for r in SC['detail']]
    sheet(wb, '1-5종II', ['계통', '번호', '수술명', '종', '비고(약관 조건)'], rows15, [14, 8, 46, 7, 56],
          '1-5종 수술분류표Ⅱ(별표76) — 수술 술식 기준 · 급여/비급여 무관 · 질병·상해 공통. 종이 높을수록 지급액이 큼. 아래 각주는 「1-5종 각주」 시트')
    notes15 = [
        ['주1', '근골(筋骨) 수술에서 발정술 등 내고정물 제거술은 제외 · 치·치은·치근·치조골 처치, 임플란트 등 치과 처치·수술도 제외'],
        ['주2', '1~87항 수술 중 내시경·카테터 등 경피적 수술은 88항 적용. 단 복강경·흉강경 수술은 해당 부위(1~87항) 수술로 적용'],
        ['시각기', '약물주입술은 제외'],
        ['88항', '검사·처치·약물주입요법은 포함하지 않음 · 60일 이내 2회 이상은 1회로 간주(이후 반복 지급 가능)'],
        ['Ⅱ 주1', '제자리암·경계성종양 수술은 「Ⅰ. 일반 질병 및 상해 치료목적의 수술」 적용'],
        ['Ⅱ 주2', '카테터를 이용한 흡인·천자·약물주입요법은 포함하지 않음'],
        ['Ⅱ 주3', '비고형암의 비관혈적 근치술은 관혈적 악성신생물 근치수술(5종)에 준해 인정 — 조혈모세포이식은 일련의 과정을 1회로 · 약물 투여용 시술(중심정맥삽관 등)만 하면 5종 아님'],
        ['Ⅱ 주4', 'Ⅱ 항목에서 복강경·흉강경 수술은 1. 또는 3. 으로 적용'],
        ['Ⅲ 주', '방사선 조사 : 60일 이내 2회 이상은 1회로 간주'],
        ['사용지침 8', '지급 제외 : 흡인 · 천자 · 신경차단 · 미용성형 목적 · 피임 목적(불임술 후 가임 목적 포함) · 검사·진단 목적(생검·복강경검사 등) · 분류표에 없는 시술'],
        ['사용지침 5', '분류표에 없는 첨단 수술 : 신의료기술평가위원회 인정 최신 수술기법은 동일 부위 수술로 봄(내시경·경피적이면 88항 / 악성은 Ⅱ-2 우선)'],
    ]
    sheet(wb, '1-5종 각주', ['구분', '약관 내용'], notes15, [14, 110])

    # ── 1-7종 ──
    x = D['tables']['34700df0f3']['text']
    T, g = [], None
    for l in x.split('\n'):
        mg = re.match(r'■\s*(\d+)\.\s*(.*)', l)
        if mg: g = (int(mg.group(1)), mg.group(2).strip())
        mm = re.match(r'\s*\[([A-Z]\d{3})\]\s*(.*?)\s*→\s*(\d)종', l)
        if mm: T.append((g, mm.group(1), mm.group(2), int(mm.group(3))))
    # 화면 표(SC.d17)의 계통·해설 · 다듬은 이름을 수술코드로 붙인다
    info = {}
    for r in d17:
        for part in r['code'].split(','):
            a, _, b = part.strip().partition('~')
            for gg, c, nm, j in T:
                if (c == a) if not b else (c[0] == a[0] and a <= c <= b): info[c] = r
    rows17 = []
    for gg, c, nm, j in T:
        r = info.get(c, {})
        single = r.get('code') == c
        rows17.append([gg[0], gg[1], c, r['name'] if single else nm, '%d종' % j, r.get('sys', ''), r.get('desc', '') if single else ''])
    sheet(wb, '1-7종', ['구분번호', '수술구분', '수술코드(ADRG)', '수술명', '종', '계통', '해설(참고)'], rows17, [8, 34, 12, 48, 7, 16, 40],
          '1-7종 수술분류표(별표3) — 건강보험 급여 수술(DRG/ADRG 코드) 기준 · 비급여 제외 · 2·3번째 자리 60~99, Error DRG(961·962·963·990·999)·960 DRG 제외. 상해·질병 담보 별도')

    # ── 131대 대상질병(약관 분류표 줄) ──
    rows131 = [[r['grp'].replace('수술비보장', '').strip(), re.sub(r'\s*\n\s*', ' ', r['dis']).strip(), re.sub(r'\s*\n\s*', ' ', r['kcd']).strip(), r['sys']] for r in d131]
    sheet(wb, '131대 대상질병', ['보장군', '대상질병(약관)', 'KCD', '신체계통'], rows131, [26, 70, 30, 18],
          '131대질병수술비 — 질병코드(KCD) 기준, 지정 질병으로 수술해야 지급(암 미보장). 130대 = 131대 − 추간판장애(M50 경추 · M51 기타). 다빈도62대 ↔ 130대는 다빈도61대')

    # ── KCD 전개(특약 마스터) ──
    def subs(pid):
        return {re.sub(r'.*\[|\]$', '', r['n']).replace('61대', '6x대').replace('62대', '6x대'): r for r in D['riders'] if r.get('parent') == pid}
    S130, S131 = subs('통123'), subs('케219')
    P130, P131 = by['통123'], by['케219']
    def cov(r, p, c): return covers(r['k'], c) and not covers(r.get('x') or [], c) and not covers(p.get('x') or [], c)
    rowsk = []
    for key, r in S131.items():
        r0 = S130.get(key)
        for i, k in enumerate(r['k']):
            lab = r['l'][i] if i < len(r.get('l') or []) else ''
            name = cn.get(k) or lab
            c0 = _expand(k.replace('-', '~'))[0]
            rowsk.append([key.replace('6x대', '61대/62대'), k, name, '보장' if cov(r, P131, c0) else '제외',
                          ('보장' if r0 and cov(r0, P130, c0) else '미보장')])
        for xc in r.get('x') or []:
            rowsk.append([key.replace('6x대', '61대/62대'), xc, (cn.get(xc) or '') + ' ← 이 세부보장에서 제외', '제외', '제외'])
    for xc in P131.get('x') or []:
        rowsk.append(['(131대·130대 전체)', xc, (cn.get(xc) or '') + ' ← 약관 「…제외」 — 어느 세부보장에도 없음', '제외', '제외'])
    sheet(wb, '130·131대 KCD', ['세부보장', 'KCD', '질병명', '131대', '130대'], rowsk, [26, 12, 60, 9, 9],
          '특약 마스터(영업지원도구 케219 · 통123) 기준 KCD 목록. 「제외」는 해당 세부보장에서 빼는 코드 — 다른 세부보장(예 : 특정31대)에서 보장될 수 있음. 130대 미보장 = M50 · M51')

    # ── 상품별 탑재 ──
    rowsp = []
    kinds = [('1-5종Ⅱ', r'1-5종'), ('1-7종', r'1-7종'), ('130대', r'130대질병수술비'), ('131대', r'131대질병수술비')]
    for r in D['riders']:
        if r.get('parent'): continue
        for kd, pat in kinds:
            if re.search(pat, r['n']) and '수술' in r['n']:
                rowsp.append([kd, r['p'], r['n'], r['id'], '영업지원도구']); break
    ex = json.load(io.open(os.path.join(ROOT, 'proposal_smart', 'db_terms_extra.json'), encoding='utf-8'))
    for r in ex['riders']:
        if '[' in r['n'] or r.get('parent'): continue
        for kd, pat in kinds:
            if re.search(pat, r['n']) and '수술' in r['n']:
                rowsp.append([kd, r['p'], r['n'], r['id'], '스마트 제안서 보강(약관 파싱)']); break
    rowsp.sort(key=lambda z: ([k for k, _ in kinds].index(z[0]), z[4], z[1], z[2]))
    sheet(wb, '상품별 탑재', ['수술비 종류', '상품', '특약명', 'id', '출처'], rowsp, [12, 22, 50, 10, 26])

    # ── 대표 질병(참고) : docs/surgery_rep_diseases.json — 코드는 WHO ICD-10(2019)에 있는 것만, 영문 공식명을 함께 싣는다 ──
    rep = os.path.join(ROOT, 'docs', 'surgery_rep_diseases.json')
    n_rep = {}
    if os.path.exists(rep):
        R = json.load(io.open(rep, encoding='utf-8'))
        def g131(code):
            """대표 질병이 130대 · 131대의 어느 세부보장에 드는가(특약 마스터 기준)"""
            hit = [k for k, r in S131.items() if cov(r, P131, code)]
            hit0 = [k for k, r in S130.items() if cov(r, P130, code)]
            a = ', '.join(h.replace('6x대', '62대') for h in hit) or '-'
            b = ', '.join(h.replace('6x대', '61대') for h in hit0) or '-'
            return a, b
        def judge7(c):
            """1-7종 [질병] 특약 제4조 ② 의 KCD 제외 · 상해는 [상해] 특약"""
            a, n = c[0], int(c[1:3])
            if a in 'ST': return '상해 담보 — 1-7종[상해]'
            if a == 'Q': return '⚠ 1-7종[질병] 미보장 — 선천기형(Q00~Q99) 약관 제외'
            if a == 'F' and n >= 4: return '⚠ 1-7종[질병] 미보장 — 정신 및 행동장애(F04~F99) 제외'
            if a == 'K' and n <= 8: return '⚠ 1-7종[질병] 미보장 — 치과질환(K00~K08) 제외'
            if a == 'P': return '⚠ 1-7종[질병] 미보장 — 출생전후기 병태(P00~P96) 제외'
            if a == 'N' and 96 <= n <= 98: return '△ 보장개시일부터 2년 지난 후 발생분만 — 습관성 유산·불임(N96~N98)'
            if a == 'Z': return '질병 담보 — 상태 코드(원인 질병으로 판단)'
            return '질병 담보 — 1-7종[질병]'
        def judge5(c):
            return '상해 담보 — 상해 1-5종 / 수술비Ⅱ(1-5종)' if c[0] in 'ST' else '질병 담보 — 질병 1-5종 / 수술비Ⅱ(1-5종)'
        rows7 = []
        for gg in sorted({(t[0][0], t[0][1]) for t in T}):
            it = R['1-7종'].get('7-%03d' % gg[0], {})
            js = sorted({t[3] for t in T if t[0][0] == gg[0]})
            cs = [t[1] for t in T if t[0][0] == gg[0]]
            for d in it.get('diseases', []):
                a, b = g131(d['kcd'])
                rows7.append([gg[0], gg[1], ', '.join(cs), '~'.join('%d종' % j for j in (js[0], js[-1])) if len(js) > 1 else '%d종' % js[0],
                              d['name'], d['kcd'], judge7(d['kcd']), d.get('who', ''), d.get('note', ''), a, b])
        sheet(wb, '1-7종 대표질병', ['구분번호', '수술구분', '수술코드', '종', '대표 질병', 'KCD', '약관상 지급 판단', 'WHO ICD-10 공식명(코드 검증)', '비고', '131대 해당 보장군', '130대 해당 보장군'],
              rows7, [8, 30, 22, 9, 34, 9, 30, 40, 26, 22, 22],
              '수술구분별로 이 수술을 받게 되는 대표 질병(참고용 · 전부가 아님). 1-7종은 질병이 아니라 「급여 수술코드(ADRG)」로 지급 — 같은 질병이라도 수술코드가 표에 있어야 함. '
              '131대/130대 칸 : 그 질병으로 수술하면 131대 · 130대 어느 보장군에도 해당하는지(특약 마스터 기준). '
              '「약관상 지급 판단」 : 1-7종[질병]은 선천기형(Q) · 정신(F04~F99) · 치과(K00~K08) · 출생전후기(P) 질병의 수술을 보상하지 않음(제4조 ②), 상해는 [상해] 특약')
        rows5 = []
        for r in SC['detail']:
            it = R['1-5종'].get('5-' + r[1], {})
            for d in it.get('diseases', []):
                a, b = g131(d['kcd'])
                rows5.append([r[0], r[1], r[2], r[3], d['name'], d['kcd'], judge5(d['kcd']), d.get('who', ''), d.get('note', ''), a, b])
        sheet(wb, '1-5종 대표질병', ['계통', '번호', '수술명', '종', '대표 질병', 'KCD', '약관상 지급 판단', 'WHO ICD-10 공식명(코드 검증)', '비고', '131대 해당 보장군', '130대 해당 보장군'],
              rows5, [12, 7, 34, 6, 34, 9, 30, 40, 30, 22, 22],
              '수술 항목별 대표 질병(참고용 · 전부가 아님). 1-5종Ⅱ는 질병코드 제외 규정이 없어 선천기형(Q)도 지급 대상(미용·피임·검사 목적만 제외) — 1-7종[질병]과 다름. 1-5종Ⅱ는 「수술 술식」으로 지급 — 질병이 달라도 같은 수술이면 같은 종. 약관에 질병별 예외가 있는 것(담석증 K80 → 2종 · 유방의 비대 N62 → 1종)은 비고에 적음')
        n_rep = {'7': len(rows7), '5': len(rows5)}

    # ── 안내 · 확인 결과 (맨 앞) ──
    ws = wb.create_sheet('안내', 0)
    ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 90
    lines = [
        ('수술비 3종 — 보장 수술 · 질병 · KCD 정리', None, None),
        ('기준일', '2026-10-20', '약관 : 케어프리 M-Basket 4.0 · 운전자 2608 3.0 · 통합간편 2607 (영업지원도구 탑재 약관)'),
        (None, None, None),
        ('시트', '건수', '내용'),
        ('1-5종II', '=COUNTA(\'1-5종II\'!C3:C200)', '수술 술식 기준 88개 항목 + 악성신생물 + 방사선 조사 (종 · 약관 조건)'),
        ('1-5종 각주', '=COUNTA(\'1-5종 각주\'!A2:A100)', '약관 각주 · 지급 제외 수술'),
        ('1-7종', '=COUNTA(\'1-7종\'!C3:C2000)', '급여 수술 ADRG 수술코드 — 266개 수술구분'),
        ('131대 대상질병', '=COUNTA(\'131대 대상질병\'!B3:B1000)', '약관 분류표의 대상질병 줄(11개 보장군)'),
        ('130·131대 KCD', '=COUNTA(\'130·131대 KCD\'!B3:B2000)', '세부보장별 KCD 전개 · 130대/131대 보장 여부'),
        ('상품별 탑재', '=COUNTA(\'상품별 탑재\'!C2:C1000)', '어느 상품에 어떤 수술비 특약이 있는지'),
    ]
    if n_rep:
        lines += [
            ('1-7종 대표질병', '=COUNTA(\'1-7종 대표질병\'!E3:E3000)', '수술구분별 대표 질병 · KCD · 131대/130대 해당 보장군 (참고용)'),
            ('1-5종 대표질병', '=COUNTA(\'1-5종 대표질병\'!E3:E3000)', '수술 항목별 대표 질병 · KCD · 131대/130대 해당 보장군 (참고용)'),
        ]
    lines += [
        (None, None, None),
        ('1-7종 종별 수술코드 수', None, None),
    ]
    for j in range(1, 8):
        lines.append(('%d종' % j, '=COUNTIF(\'1-7종\'!E3:E2000,"%d종")' % j, ''))
    lines += [
        (None, None, None),
        ('확인 결과(2026-10-20)', None, None),
        ('1-7종', '일치', '수술코드 630개 · 종 — 약관 PDF 4권(케어프리 · 운전자 · 통합간편 일반/통합간편심사형)에서 직접 읽은 값과 모두 같음. 화면 이름 중 조건이 빠진 67줄을 약관 이름으로(「치료적」 자궁소파술 · 「재원기간 2일 이상」 등)'),
        ('1-5종Ⅱ', '일치', '99줄 종 모두 약관과 같음. 상품별 분류표 글도 같음(줄바꿈 차이뿐). 빠져 있던 약관 각주를 덧붙임'),
        ('131대 · 130대', '1건 바로잡음', '황반변성(H35.3) : 약관 다빈도61대·62대 분류표에 「시각질환 · 황반변성 · 황반 및 후극부의 변성 H35.3」 독립 줄이 있어 보장. 10-20 오전(v8.78)에 「H30-H36 단, H35.3 제외」 문구만 보고 제외로 바꾼 것을 되돌림 — 130대 · 131대 11개 상품'),
        ('131대 · 130대', '일치', '세부보장 11개 코드 = 화면 표 = 약관. 130대와 131대 차이는 추간판장애(M50 · M51)뿐. 제외 코드 K23.0 · M74 · M78 · N39.3 · N39.4 는 어느 세부보장에도 없음(약관대로 미보장)'),
        ('주의', '', 'H35.0(배경망막병증)은 코드 하나에 고혈압성(특정31대)과 그 밖(다빈도)이 섞여 있어 코드만으로 나눌 수 없다 — 지급 판단 시 상병명 확인'),
        ('대표 질병 시트', '참고용', '대표 질병은 약관에 적힌 목록이 아니라 「이 수술을 흔히 받게 되는 질병」 예시다(의학 자료 · 건강보험심사평가원 KDRG 분류 · 의료기관 안내 확인). '
                                   '질병코드는 모두 WHO ICD-10(2019) 목록에 실제로 있는 코드만 싣고 영문 공식명을 옆에 붙여 대조할 수 있게 했다. 실제 지급은 수술코드(1-7종) · 수술 술식(1-5종)과 약관 조건으로 판단'),
        ('출처', '', '영업지원도구(tool.html) 특약 마스터 · 수술비 비교 표 · 약관 별표 원문 · docs/surgery_rep_diseases.json. 생성 : scripts/build_surgery_xlsx.py'),
    ]
    for i, (a, b, c) in enumerate(lines, 1):
        for j, v in enumerate((a, b, c), 1):
            if v is None: continue
            cell = ws.cell(i, j, v); cell.font = Font(name=FONT, size=10); cell.alignment = Alignment(wrap_text=True, vertical='top')
        if a and b is None and c is None:
            ws.cell(i, 1).font = Font(name=FONT, size=13 if i == 1 else 11, bold=True, color='0B7C6F')
        if a in ('시트',):
            for j in range(1, 4):
                ws.cell(i, j).font = Font(name=FONT, bold=True, color='FFFFFF'); ws.cell(i, j).fill = HEAD
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(OUT, '· 1-5종II %d · 1-7종 %d · 131대 %d · KCD %d · 탑재 %d' % (len(rows15), len(rows17), len(rows131), len(rowsk), len(rowsp)))


if __name__ == '__main__':
    main()
