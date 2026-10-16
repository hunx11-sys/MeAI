# -*- coding: utf-8 -*-
"""
간편보험 예외질환 심사 매뉴얼 만들기 — 엑셀(경우의 수) + MeAI RAG 용 MD

원천은 영업지원도구 상품 라인업(products.html)의 「예외질환 판정」 탭 데이터(PLANS · ITEMS · pNew) 하나뿐이다.
판정 규칙도 화면(judgeOne · judgePlan)과 똑같이 옮겨, 엑셀·MD 의 경우의 수가 화면 결과와 어긋나지 않게 한다.
예외질환 기준이 바뀌면 products.html 만 고치고 이 스크립트를 다시 돌린다.

실행 : python3 scripts/build_exception_manual.py
산출 : dist/simple_exception_disease_2610.xlsx
       docs/meai_rag/simple_exception_disease_manual_2610.md
필요 : node (products.html 안의 자바스크립트 데이터를 그대로 읽는다) · openpyxl
"""
import json, os, re, subprocess, sys, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'products.html')
XLSX = os.path.join(ROOT, 'dist', 'simple_exception_disease_2610.xlsx')
MD = os.path.join(ROOT, 'docs', 'meai_rag', 'simple_exception_disease_manual_2610.md')

FORM = {'I': '입원만 인정(수술·시술하면 불가)', 'IS': '입원 또는 수술·시술 모두 인정', 'S': '수술·시술한 경우만 인정'}
FORM_S = {'I': '입원만', 'IS': '입원·수술', 'S': '수술만'}


def load():
    s = open(SRC, encoding='utf-8').read()
    a, b = s.index('const PLANS=['), s.index('/* =============== 2026.10.16 시행')
    c, d = s.index('const EX_NEW_DATE'), s.index('/* =============== 초성 / 검색 유틸')
    js = s[a:b] + s[c:d] + '\nconsole.log(JSON.stringify({PLANS,ITEMS,EX_NEW_DATE,EX_SURG_RIDERS}));'
    with tempfile.NamedTemporaryFile('w', suffix='.js', delete=False, encoding='utf-8') as f:
        f.write(js)
    out = subprocess.run(['node', f.name], capture_output=True, text=True, check=True).stdout
    os.unlink(f.name)
    ver = re.search(r'간편심사 예외질환 · (ver\d+)', s).group(1)
    return json.loads(out), ver


# ── 화면과 같은 규칙 (products.html judgeOne) ──
def rv(it, pid, rule):
    if rule == 'NEW' and it.get('pNew') and it['pNew'].get(pid):
        return it['pNew'][pid]
    return it['p'].get(pid)


def max_n(v):
    if v.get('maxN'):
        return v['maxN']
    m = re.search(r'(\d+)\s*회', v.get('lim') or '')
    return int(m.group(1)) if m else 0


def judge(it, plan, rule, op, adm=0, out=0, cnt=1, c5=1, c10=0):
    v = rv(it, plan['id'], rule)
    if not v:
        return '불가', '예외질환 목록에 없음'
    if op and v['c'] == 'I':
        return '불가', '수술하면 불가, 입원만 인정'
    if not op and v['c'] == 'S':
        return '불가', '수술한 경우만 인정'
    days = adm + out if plan['id'] == 'SIL' else adm
    if days > v['d']:
        return '불가', '%s %d일, 한도 %d일 초과' % (plan['unit'], days, v['d'])
    if op and it.get('colon'):
        tot = c5 + c10
        if v.get('colon') == 'new':
            if c5 > 2: return '불가', '5년 이내 %d회, 최대 2회 초과' % c5
            if tot > 3: return '불가', '10년 내 %d회, 최대 3회 초과' % tot
        elif tot > 2:
            return '불가', '수술 %d회, 2회 초과' % tot
    elif op:
        n = max_n(v)
        if n and cnt > n:
            return '불가', '%s %d회, 한도 %d회 초과' % (it.get('cntKind') or '수술', cnt, n)
    if not op and v.get('noop'):
        return '조건부 가능', '%s %d일 이내 · %s' % (plan['unit'], v['d'], v['noop'])
    return '가능', '%s %d일 이내%s' % (plan['unit'], v['d'], (', ' + v['lim']) if v.get('lim') else '')


def scenarios(it, plan, rule):
    """질환 × 설계유형 하나의 경우의 수 : (치료이력 설명, 결과, 사유)"""
    v = rv(it, plan['id'], rule)
    if not v:
        return [('어떤 치료든', '불가', '이 설계유형의 예외질환 목록에 없음(그 外 → 인수불가)')]
    d, sil, out = v['d'], plan['id'] == 'SIL', []
    unit = '입·통원 합산' if sil else '입원'
    k = it.get('cntKind') or '수술'

    def add(desc, **kw):
        out.append((desc,) + judge(it, plan, rule, **kw))
    if v['c'] in ('I', 'IS'):
        add('입원만(비수술) · %s %d일' % (unit, d), op=False, adm=d)
        add('입원만(비수술) · %s %d일' % (unit, d + 1), op=False, adm=d + 1)
    if v['c'] == 'S':
        add('입원만(비수술) · %s %d일' % (unit, min(d, 3)), op=False, adm=min(d, 3))
    if v['c'] in ('S', 'IS'):
        if it.get('colon'):
            add('수술 · %s %d일 · 5년 이내 1회' % (unit, d), op=True, adm=d, c5=1, c10=0)
            add('수술 · %s %d일 · 5년 이내 1회' % (unit, d + 1), op=True, adm=d + 1, c5=1, c10=0)
            for c5, c10 in ((2, 0), (2, 1), (1, 2), (0, 3), (3, 0), (2, 2), (1, 3)):
                add('수술 · %s %d일 · 5년 이내 %d회 + 5~10년 전 %d회' % (unit, d, c5, c10), op=True, adm=d, c5=c5, c10=c10)
        else:
            add('수술 · %s %d일 · %s 1회' % (unit, d, k), op=True, adm=d, cnt=1)
            add('수술 · %s %d일 · %s 1회' % (unit, d + 1, k), op=True, adm=d + 1, cnt=1)
            n = max_n(v)
            if n:
                if n > 1:
                    add('수술 · %s %d일 · %s %d회' % (unit, d, k, n), op=True, adm=d, cnt=n)
                add('수술 · %s %d일 · %s %d회' % (unit, d, k, n + 1), op=True, adm=d, cnt=n + 1)
    if v['c'] == 'I':
        add('수술·시술 · %s %d일' % (unit, min(d, 3)), op=True, adm=min(d, 3))
    if sil and v['c'] in ('I', 'IS'):
        a = max(1, d // 2)
        add('입원만 · 입원 %d일 + 통원 %d일 (합산 %d일)' % (a, d - a + 1, d + 1), op=False, adm=a, out=d - a + 1)
    return out


def plabel(p):
    return p['nm'] if p['grp'] == '유병자실손' else '%s %s' % (p['grp'], p['nm'])


def cond_text(v):
    if not v:
        return ''
    t = [v.get('lim') or '', v.get('x') or '', v.get('noop') or '']
    return ' · '.join(x for x in t if x)


# ═══════════════ 엑셀 ═══════════════
def build_xlsx(D, ver):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    P, I = D['PLANS'], D['ITEMS']
    wb = Workbook()
    H = PatternFill('solid', fgColor='1F3A5F'); HF = Font(bold=True, color='FFFFFF')
    OK = PatternFill('solid', fgColor='E3F6EC'); NO = PatternFill('solid', fgColor='FBE4E6')
    WN = PatternFill('solid', fgColor='FFF4D6'); CH = PatternFill('solid', fgColor='FFF1CC')
    thin = Side(style='thin', color='D0D7E2'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top')

    def sheet(ws, head, rows, widths, fills=None):
        ws.append(head)
        for c in ws[1]:
            c.fill, c.font, c.alignment, c.border = H, HF, Alignment(wrap_text=True, vertical='center', horizontal='center'), B
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment, c.border = WR, B
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'B2' if len(head) > 6 else 'A2'
        ws.auto_filter.ref = ws.dimensions
        if fills:
            fills(ws)

    def color_result(col):
        def f(ws):
            for row in ws.iter_rows(min_row=2):
                c = row[col]
                c.fill = OK if c.value == '가능' else WN if c.value == '조건부 가능' else NO if c.value == '불가' else c.fill
        return f

    # 0. 안내
    ws = wb.active; ws.title = '안내'
    lines = [
        ['간편보험 예외질환 심사 매뉴얼 — 경우의 수 (%s · %s 시행 기준 반영)' % (ver, D['EX_NEW_DATE'])],
        [''],
        ['원천', '영업지원도구 > 상품 라인업 > 「예외질환 판정」 탭 데이터(출처 : 간편상품(SI) 예외질환 안내 화면 5종 · 예외질환 인정기준 변경 안내 시행 2026.10.16)'],
        ['설계유형', ' · '.join('%s %s(최대 %d개 · %s 일수)' % (p['grp'], p['nm'], p['max'], p['unit']) for p in P)],
        ['치료형태', '입원만 = 수술·시술하면 불가 / 입원·수술 = 둘 다 인정 / 수술만 = 수술·시술한 경우만 인정'],
        ['일수', '간편상품·간편정기는 입원일수만, 유병자실손은 입원+통원 일수를 합산해 한도와 비교'],
        ['횟수', '목록에 「○회 이내」가 있으면 수술(수면무호흡은 입원) 횟수가 그 이하여야 함. 대장용종(신기준)은 5년 이내 최대 2회 · 10년 내 최대 3회'],
        ['시트', '인수 필수조건 · 질환×설계유형(신기준) · 질환×설계유형(종전) · 2026.10.16 변경 · 경우의 수(신기준) · 여러 질환 · 판정 순서'],
        ['주의', '사전 참고용입니다. 최종 인수 여부는 청약 심사 결과에 따릅니다. 질병코드는 검색 편의용 참고값이며 인수 판단 기준은 질환명입니다.'],
        ['다시 만들기', 'python3 scripts/build_exception_manual.py (products.html 의 판정 데이터·규칙을 그대로 읽음)'],
    ]
    for l in lines:
        ws.append(l)
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 120
    for row in ws.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        for c in row: c.alignment = WR

    # 1. 인수 필수조건
    ws = wb.create_sheet('인수 필수조건')
    rows = [
        ['간편상품(SI)·간편정기', '1', '예외질환(그 外) & 최종 치료종료 7일 경과 & 재발·합병증·후유장해 없음'],
        ['간편상품(SI)·간편정기', '2', '예외질환은 최대 5개까지 인수, 6개 이상이면 사고다발로 인수불가'],
        ['간편상품(SI)·간편정기', '3', '치료방법(입원·수술 등)과 치료기간(입원 ○일 이내 등)을 명기한 범위 안에서만 인수. 다른 치료방법이거나 기간이 넘으면 인수불가'],
        ['간편상품(SI)·간편정기', '4', '예외질환은 경증 위험에만 적용 — 기준에 해당해도 고위험 RISK(중대질환·유의고객 등)가 있으면 인수 제한 가능'],
        ['유병자 실손의료비', '1', '예외질환(그 外) & 최종 치료종료 3개월 경과 & 재발·합병증·후유장해 없음'],
        ['유병자 실손의료비', '2', '예외질환은 최대 3개까지 인수, 4개 이상이면 사고다발로 인수불가'],
        ['유병자 실손의료비', '3', '치료방법과 치료기간을 명기한 범위 안에서만 인수. 일수는 입원과 통원을 합산'],
        ['유병자 실손의료비', '4', '고위험 RISK 요인이 있으면 인수 제한 가능'],
        ['공통', '그 外', '목록에 없는 질환으로 입원(실손은 입·통원)하거나 수술했다면 예외질환이 아니므로 인수불가'],
        ['공통', '대안', '5개 설계유형 모두 막히면 예외질환을 운영하지 않는 355 수술고지형을 검토'],
    ]
    sheet(ws, ['대상', '번호', '조건'], rows, [24, 8, 110])

    # 2·3. 질환 × 설계유형
    def matrix(rule, title):
        ws = wb.create_sheet(title)
        head = ['ID', '질환', '세부 질환', '질병코드(참고)', '검색어']
        for p in P:
            head += ['%s\n인정' % p['nm'], '%s\n치료형태' % p['nm'], '%s\n일수 한도' % p['nm'], '%s\n조건·횟수' % p['nm']]
        rows = []
        for it in I:
            r = [it['id'], it['n'], it.get('s', ''), it.get('code', ''), it.get('kw', '')]
            for p in P:
                v = rv(it, p['id'], rule)
                if v:
                    r += ['O', FORM_S[v['c']], '%s %d일' % (p['unit'], v['d']), cond_text(v)]
                else:
                    r += ['X', '', '', '']
            rows.append(r)

        def fill(ws):
            for row in ws.iter_rows(min_row=2):
                for j in range(5, len(row), 4):
                    row[j].fill = OK if row[j].value == 'O' else NO
                    row[j].alignment = Alignment(horizontal='center', vertical='top')
                it = next(x for x in I if x['id'] == row[0].value)
                if rule == 'NEW' and it.get('pNew'):
                    for j in range(5, len(row), 4):
                        pid = P[(j - 5) // 4]['id']
                        if it['pNew'].get(pid):
                            for c in row[j:j + 4]: c.fill = CH
        sheet(ws, head, rows, [6, 28, 34, 20, 30] + [6, 10, 11, 30] * len(P), fill)
    matrix('NEW', '질환×설계유형(신기준)')
    matrix('OLD', '질환×설계유형(종전)')

    # 4. 변경
    ws = wb.create_sheet('2026.10.16 변경')
    rows = []
    for it in I:
        if not it.get('pNew'):
            continue
        for p in P:
            n = it['pNew'].get(p['id'])
            if not n:
                continue
            o = it['p'].get(p['id'])
            ot = '%s · %s %d일 · %s' % (FORM_S[o['c']], p['unit'], o['d'], cond_text(o)) if o else '없음'
            nt = '%s · %s %d일 · %s' % (FORM_S[n['c']], p['unit'], n['d'], cond_text(n))
            rows.append([it['n'], plabel(p), ot, nt])
    rows.append(['(질병)수술비 범위', '요로결석 비수술 인수 시 가입 불가', '', D['EX_SURG_RIDERS']])
    sheet(ws, ['예외질환', '설계유형', '종전(~10.15)', '신기준(10.16~)'], rows, [30, 24, 50, 60])

    # 5. 경우의 수
    ws = wb.create_sheet('경우의 수(신기준)')
    rows = []
    for it in I:
        for p in P:
            for desc, res, why in scenarios(it, p, 'NEW'):
                rows.append([it['id'], it['n'], plabel(p), desc, res, why])
    sheet(ws, ['ID', '질환', '설계유형', '치료이력(경우)', '판정', '사유·조건'], rows, [6, 30, 22, 46, 11, 50], color_result(4))
    n_scen = len(rows)

    # 6. 여러 질환
    ws = wb.create_sheet('여러 질환')
    rows = []
    for p in P:
        for n in range(1, p['max'] + 2):
            ok = n <= p['max']
            rows.append([plabel(p), n, '가능' if ok else '불가',
                         '각 질환이 모두 기준 안이면 인수' if ok else '%d개로 한도 %d개 초과 → 사고다발 인수불가' % (n, p['max'])])
    sheet(ws, ['설계유형', '예외질환 개수', '판정', '설명'], rows, [26, 14, 10, 60], color_result(2))

    # 7. 판정 순서
    ws = wb.create_sheet('판정 순서')
    steps = [
        ['1', '목록 확인', '질환이 그 설계유형의 예외질환 목록에 있는가? 없으면 인수불가(그 外)'],
        ['2', '치료형태', '수술·시술했는데 「입원만」 질환이면 불가 / 수술 안 했는데 「수술만」 질환이면 불가'],
        ['3', '일수', '입원일수(유병자실손은 입원+통원 합산)가 한도 이하인가'],
        ['4', '횟수', '수술(수면무호흡은 입원) 횟수가 「○회 이내」 이하인가. 대장용종(신기준)은 5년 이내 ≤2회 · 10년 내 ≤3회'],
        ['5', '조건부', '요로결석(신기준)을 수술 없이 치료했으면 인수 가능하나 (질병)수술비 가입 불가 → 「조건부 가능」'],
        ['6', '개수', '고른 예외질환이 간편상품·간편정기 5개 / 유병자실손 3개 이하인가'],
        ['7', '공통 조건', '최종 치료종료 7일(실손 3개월) 경과 · 재발·합병증·후유장해 없음 · 고위험 RISK 없음 — 도구는 확인하지 않으므로 사람이 확인'],
    ]
    sheet(ws, ['순서', '확인', '내용'], steps, [8, 14, 110])
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    return n_scen


# ═══════════════ MD (RAG) ═══════════════
def build_md(D, ver):
    P, I = D['PLANS'], D['ITEMS']
    L = []
    w = L.append
    w('---')
    w('title: 간편보험 예외질환 심사 매뉴얼 (메리츠화재 · %s)' % ver)
    w('doc_type: underwriting_manual')
    w('effective_date: %s' % D['EX_NEW_DATE'])
    w('source: 영업지원도구 상품 라인업 「예외질환 판정」 데이터 · 간편상품(SI) 예외질환 안내 화면 5종 · 예외질환 인정기준 변경 안내(시행 2026.10.16)')
    w('scope: 간편상품(일반설계·3대진단플랜·어깨/무릎/디스크플랜) · 간편정기(정기·상조) · 유병자 실손의료비')
    w('notice: 사전 참고용. 최종 인수 여부는 청약 심사 결과에 따른다. 질병코드는 검색 참고값이며 판단 기준은 질환명이다.')
    w('---')
    w('')
    w('# 간편보험 예외질환 심사 매뉴얼')
    w('')
    w('> 이 문서는 MeAI 답변용 근거 문서다. 질문에 답할 때는 아래 「판정 순서」대로 확인하고, 질환별 절의 표에 있는 값만 쓴다. '
      '표에 없는 질환·치료방법·일수는 추정하지 말고 「예외질환 목록에 없음 → 인수불가(그 外)」 또는 「심사 문의」로 답한다. '
      '%s 이후 청약은 「신기준」, 그 이전은 「종전 기준」을 쓴다.' % D['EX_NEW_DATE'])
    w('')
    w('## 1. 용어')
    w('')
    w('- **예외질환** : 간편보험 알릴의무(고지)에서 입원·수술 이력이 있어도 정해진 치료형태·기간 안이면 인수하는 경증 질환 목록.')
    w('- **설계유형 5개** : ' + ' · '.join(plabel(p) for p in P) + '.')
    w('- **치료형태** : 「입원만」= 입원만 인정(수술·시술하면 불가) / 「입원·수술」= 둘 다 인정 / 「수술만」= 수술·시술한 경우만 인정.')
    w('- **일수 한도** : 간편상품·간편정기는 입원일수, 유병자실손은 입원+통원 일수 합산.')
    w('- **그 外** : 목록에 없는 질환. 이 질환으로 입원(실손은 입·통원)하거나 수술했으면 인수불가.')
    w('- **(질병)수술비** : ' + D['EX_SURG_RIDERS'] + '.')
    w('')
    w('## 2. 인수 필수조건')
    w('')
    w('| 대상 | 치료종료 경과 | 예외질환 최대 개수 | 초과 시 | 일수 기준 |')
    w('|---|---|---|---|---|')
    w('| 간편상품(SI) 3개 플랜 · 간편정기 | 최종 치료종료 7일 | 5개 | 6개 이상 사고다발 인수불가 | 입원일수 |')
    w('| 유병자 실손의료비 | 최종 치료종료 3개월 | 3개 | 4개 이상 사고다발 인수불가 | 입원+통원 합산 |')
    w('')
    w('공통 : 재발·합병증·후유장해가 없어야 하고, 명기된 치료방법·기간 범위 안에서만 인수한다. 기준 안이어도 고위험 RISK(중대질환·유의고객 등)가 있으면 인수가 제한될 수 있다.')
    w('')
    w('## 3. 판정 순서')
    w('')
    w('1. **목록** — 그 설계유형의 예외질환 목록에 있는가? 없으면 인수불가.')
    w('2. **치료형태** — 수술했는데 「입원만」이면 불가, 수술 안 했는데 「수술만」이면 불가.')
    w('3. **일수** — 한도 이하인가(실손은 입·통원 합산).')
    w('4. **횟수** — 「○회 이내」면 수술 횟수(수면무호흡은 입원 횟수)가 그 이하인가. 대장용종(신기준)은 5년 이내 최대 2회 · 10년 내 최대 3회.')
    w('5. **조건부** — 요로결석(신기준)을 수술 없이 치료했으면 인수 가능, 단 (질병)수술비 가입 불가.')
    w('6. **개수** — 간편상품·간편정기 5개, 유병자실손 3개 이하인가.')
    w('7. **공통 조건** — 치료종료 경과기간 · 재발·합병증·후유장해 · 고위험 RISK.')
    w('')
    w('5개 설계유형이 모두 막히면 예외질환을 운영하지 않는 **355 수술고지형**을 검토한다. 일반설계로 안 되면 3대진단플랜·어깨/무릎/디스크플랜·간편정기에서 되는지 함께 본다(질환마다 인정 설계유형이 다르다).')
    w('')
    w('## 4. 2026.10.16 시행 변경 (유병력자실손 제외 간편상품)')
    w('')
    w('| 예외질환 | 종전(~10.15) | 신기준(10.16~) |')
    w('|---|---|---|')
    w('| 대장용종(3.6.5~3.10.5) | 수술 2회 이내 | 10년 내 3회, 단 5년 이내 최대 2회 (5년 이내·5~10년 = 2·1 / 1·2 / 0·3회) |')
    w('| 수면무호흡 | 입원 1회, 수술 필수 | 입원 2회, 수술 필수 |')
    w('| 요로결석 | 수술 필수, 수술 2회 이내 | 수술 시 2회 이내 · 비수술도 인정(단 (질병)수술비 가입 불가) |')
    w('| 화상 | 표재성, 비수술 | 표재성·심재성 무관, 수술 무관 · 수술 3회 이내 |')
    w('| 자궁내막증 | 수술 1회(양측 난소절제술 필수) | 수술 1회(수술 종류 무관) |')
    w('')
    w('유병자 실손의료비는 위 변경이 적용되지 않고 종전 기준 그대로다.')
    w('')
    w('## 5. 질환별 기준과 경우의 수')
    w('')
    for it in I:
        w('### %s %s' % (it['id'], it['n']))
        w('')
        meta = []
        if it.get('s'): meta.append('- 세부 질환 : ' + it['s'])
        if it.get('code'): meta.append('- 질병코드(참고) : ' + it['code'])
        if it.get('kw'): meta.append('- 다른 이름·검색어 : ' + it['kw'])
        meta.append('- 분류 : ' + it.get('g', ''))
        L.extend(meta)
        w('')
        rules = ['NEW', 'OLD'] if it.get('pNew') else ['NEW']
        for rule in rules:
            if len(rules) == 2:
                w('**%s**' % ('신기준(2026.10.16~)' if rule == 'NEW' else '종전 기준(~2026.10.15)'))
                w('')
            w('| 설계유형 | 인정 | 치료형태 | 일수 한도 | 조건·횟수 |')
            w('|---|---|---|---|---|')
            for p in P:
                v = rv(it, p['id'], rule)
                if v:
                    w('| %s | O | %s | %s %d일 | %s |' % (plabel(p), FORM_S[v['c']], p['unit'], v['d'], cond_text(v) or '-'))
                else:
                    w('| %s | X | - | - | 목록에 없음(인수불가) |' % plabel(p))
            w('')
        w('경우의 수(신기준) :')
        w('')
        seen = {}
        for p in P:
            for desc, res, why in scenarios(it, p, 'NEW'):
                seen.setdefault((desc, res, why), []).append(p['nm'])
        for (desc, res, why), pls in seen.items():
            w('- %s → **%s** (%s) — %s' % (desc, res, why, ' · '.join(pls)))
        w('')
    w('## 6. 자주 묻는 질문')
    w('')
    faq = [
        ('목록에 없는 질환으로 3일 입원했는데 간편보험 가입되나요?', '예외질환이 아니므로(그 外) 5개 설계유형 모두 인수불가다. 355 수술고지형을 검토한다.'),
        ('예외질환이 여러 개면?', '간편상품·간편정기는 5개까지, 유병자실손은 3개까지. 각 질환이 모두 기준 안이어야 하며 개수를 넘으면 사고다발로 인수불가.'),
        ('유병자실손 일수는 어떻게 세나요?', '입원일수와 통원일수를 합산해 한도와 비교한다. 간편상품·간편정기는 입원일수만 본다.'),
        ('요로결석을 약물·체외충격파 없이 입원 치료만 했어요(10.16 이후).', '간편상품 3개 플랜·간편정기는 입원 10일 이내면 인수 가능하나 (질병)수술비는 가입할 수 없다. 유병자실손은 수술한 경우만 인정이라 불가.'),
        ('대장용종을 5년 안에 2번, 7년 전에 1번 뗐어요(10.16 이후).', '5년 이내 2회 · 10년 내 3회로 간편상품·간편정기 인수 가능(입원 5일 이내 · 현재 보유 없음). 유병자실손은 2회 이내라 불가.'),
        ('화상으로 피부이식 수술을 2번 했어요(10.16 이후).', '간편상품·간편정기는 화상 깊이·수술 무관, 수술 3회 이내라 인정(구축·강직·후유장해 없을 것). 유병자실손은 입원만 인정이라 수술하면 불가.'),
        ('기준에 맞으면 무조건 가입되나요?', '아니다. 최종 치료종료 경과기간·재발·합병증·후유장해·고위험 RISK를 함께 보며, 최종 인수는 청약 심사 결과에 따른다.'),
    ]
    for q, a in faq:
        w('**Q. %s**' % q)
        w('')
        w('A. %s' % a)
        w('')
    w('---')
    w('')
    w('이 문서는 `scripts/build_exception_manual.py` 가 영업지원도구의 판정 데이터(products.html)에서 자동으로 만든다. 경우의 수는 판정 화면과 같은 규칙으로 계산했다. 제작 : 세일즈혁신TF')
    os.makedirs(os.path.dirname(MD), exist_ok=True)
    open(MD, 'w', encoding='utf-8').write('\n'.join(L) + '\n')
    return len(L)


if __name__ == '__main__':
    D, ver = load()
    n = build_xlsx(D, ver)
    m = build_md(D, ver)
    print('질환 %d · 설계유형 %d · 변경 질환 %d' % (len(D['ITEMS']), len(D['PLANS']), sum(1 for i in D['ITEMS'] if i.get('pNew'))))
    print(os.path.relpath(XLSX, ROOT), '· 경우의 수', n, '행')
    print(os.path.relpath(MD, ROOT), '·', m, '줄 ·', os.path.getsize(MD) // 1024, 'KB')
