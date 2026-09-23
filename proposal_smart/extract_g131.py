# -*- coding: utf-8 -*-
"""'131대 수술비 질병코드 정리' 엑셀 → g131.json (그룹 → KCD 목록).
   사용 : python extract_g131.py 131대질병_질병코드정리.xlsx
   대분류(심장질환·뇌혈관질환·특정31대질병·다빈도64대질병·특정다빈도29대질병·후각특정질환·관절염․생식기질환·치핵·백내장)를 그룹으로,
   코드 범위(I00~I02)는 그대로 두고(code_hit 이 범위표기를 지원), 제외 코드는 '!코드' 로 넣는다. 소분류별 목록은 sub 에 함께 남긴다."""
import json, os, re, sys
def norm_group(g):
    g = re.sub(r'\s+', '', g or ''); g = re.sub(r'\(.*?\)', '', g)
    return g.replace('수술비보장', '')
def parse_codes(s):
    s = (s or '').replace('∼', '~').replace('－', '-')
    inc, exc = [], []
    for m in re.finditer(r'\(([^()]*?)제외\)', s):
        exc += re.findall(r'[A-Z]\d{2}(?:\.\d+)?', m.group(1))
    body = re.sub(r'\([^()]*?제외\)', ' ', s)
    for tok in re.split(r'[,\s]+', body):
        tok = tok.strip('†*+ ')
        if not tok: continue
        m = re.fullmatch(r'([A-Z]\d{2})(?:\.\d+)?[~-]([A-Z]\d{2})(?:\.\d+)?', tok)
        if m: inc.append(f'{m.group(1)}~{m.group(2)}'); continue
        m = re.fullmatch(r'[A-Z]\d{2}(?:\.\d+)?', tok)
        if m: inc.append(tok)
    return inc, exc
def extract(xlsx):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, read_only=True, data_only=True).worksheets[0]
    cur = [None, None, None]; groups = {}; subs = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 5: continue
        _, a, b, c, d, e = (list(row) + [None] * 6)[:6]
        if a: cur[0] = str(a)
        if b: cur[1] = str(b)
        if c: cur[2] = str(c)
        if not e: continue
        g = norm_group(cur[0]); sub = re.sub(r'\s+', ' ', str(cur[2] or '')).replace('수술비보장', '').strip()
        inc, exc = parse_codes(str(e))
        lst = groups.setdefault(g, [])
        for x in inc + ['!' + x for x in exc]:
            if x not in lst: lst.append(x)
        subs.setdefault(g, {}).setdefault(sub, [])
        for x in inc:
            if x not in subs[g][sub]: subs[g][sub].append(x)
    return groups, subs
if __name__ == '__main__':
    groups, subs = extract(sys.argv[1])
    base = os.path.dirname(os.path.abspath(__file__))
    json.dump(groups, open(os.path.join(base, 'g131.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    json.dump(subs, open(os.path.join(base, 'g131_sub.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    print('g131.json 생성 :', {k: len(v) for k, v in groups.items()}, '| 소분류', sum(len(v) for v in subs.values()))
