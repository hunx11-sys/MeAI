# -*- coding: utf-8 -*-
"""보험금 지급 규칙표(rules.json)를 사람이 눈으로 검수할 엑셀로 내보낸다.

    python export_rules_xlsx.py [내보낼파일.xlsx]

규칙 76줄을 위에서부터 그대로 한 줄씩 옮기고, 코드에서 쓰는 약어를 한글로 풀어 적는다.
특약 마스터(db.json) 1,757건을 실제로 이 규칙표에 통과시켜 규칙마다 몇 건이 걸리는지,
어떤 담보가 걸리는지도 함께 적는다. 규칙 내용은 손대지 않고 읽어서 옮기기만 한다.
"""
import json, os, sys, collections

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, 'out', 'rules-76.xlsx')

FONT = '맑은 고딕'
NAVY = '1A2340'
RED = 'C12027'
HEAD = PatternFill('solid', fgColor=NAVY)
BAND = PatternFill('solid', fgColor='FBF8F5')
WARN = PatternFill('solid', fgColor='FDF0F0')
FILLIN = PatternFill('solid', fgColor='FFFBE6')
THIN = Side(style='thin', color='D9D9D9')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ── 코드에서 쓰는 약어 → 한글 ─────────────────────────────────────────
KIND = {
    'dx':   '진단비',
    'surg': '수술비',
    'day':  '입원·통원 일당',
    'tx':   '치료비 (치료행위)',
    'care': '간병',
    'life': '사망·후유장해',
    'ls':   '생활지원금',
    'nonmed': '의료사례 비대상 (운전자·재물·법률·배상·치아)',
    'skip': '계산 제외 (제도성·조건부)',
    'itc_unknown': '계산 제외 + 사유 기록',
}
KCD = {
    'major':        '암 — 유사암 제외 (약관 별표 암분류로 판정)',
    'sim':          '유사암 — 제자리암·경계성종양·갑상선암·기타피부암',
    'major_or_sim': '암 또는 유사암',
    'codes':        '특약 마스터(db.json)에 실린 약관 KCD 목록과 대조',
    'codes_listed': '마스터에 약관 KCD 목록이 있으면 대조 (없으면 조건 없이 계산하고 로그)',
    'g131':         '131대질병 그룹표(g131.json)와 대조',
    'group':        '질병코드 그룹표(rules.json)와 대조',
    'hc':           '약관 별표 진료행위(수가)코드와 대조',
    'none':         '질병코드 조건 없음 — 담보명·치료행위로만 판정',
}
FREQ = {'once': '최초 1회', 'year': '연 1회', 'each': '받을 때마다'}
# 규칙에 지급 횟수를 안 적으면 보상 유형별 기본값이 쓰인다(scen_engine.py 의 h_* 기본값)
FREQ_DEFAULT = {'dx': 'once', 'surg': 'each', 'day': 'each', 'tx': 'year'}
# 사례 금액을 계산하는 보상 유형 — 나머지는 담보를 분류만 하고 금액은 만들지 않는다
CALC = ('dx', 'surg', 'day', 'tx')
MODE = {'day': '입원 1일당', 'visit': '통원 1회당', 'icu': '중환자실 1일당', 'daycare': '낮병동 입원 1일당'}
GRADE = {'1-5': '1-5종 수술분류표Ⅱ (약관 별표76)', '1-7': '1-7종 수술분류표 (약관 별표3)'}

ACTS = {
    'surg': '수술', 'robot': '다빈치 로봇수술', 'anes6': '6시간 이상 전신마취 수술',
    'chemo': '항암 화학약물치료(급여)', 'target': '표적항암약물치료',
    'immune': '면역항암약물치료', 'hormone': '항암 호르몬약물치료',
    'rad': '항암 방사선치료', 'imrt': '세기조절방사선치료(IMRT)',
    'proton': '양성자 방사선치료', 'carbon': '중입자 방사선치료',
    'thromb': '혈전용해치료(약물)', 'thrombectomy': '기계적 혈전제거술',
    'rehab': '재활치료', 'pain': '통증치료(신경차단 등)',
    'er': '응급실 진료', 'trauma_center': '권역외상센터 진료',
    'icu': '중환자실 입원', 'ecmo': '에크모(체외막산소공급)', 'crrt': '지속적 신대체요법',
    'vent': '인공호흡기', 'hypo': '저체온 치료',
    'x_ct': 'CT 검사', 'x_mri': 'MRI 검사', 'x_pet': 'PET 검사',
    'x_endo': '내시경 검사', 'x_bio': '조직검사(생검)', 'x_us': '초음파 검사',
    'x_gene': '유전자 검사', 'x_ngs': 'NGS 유전자패널검사',
    'x_punct': '천자(穿刺) 검사', 'x_bm': '골수 검사',
}
FAM = {'암': '암', '유사암': '유사암', '뇌': '뇌질환', '심장': '심장질환', '상해': '상해', '치매': '치매'}


def acts_ko(lst):
    return '\n'.join('%s  (%s)' % (ACTS.get(a, '※ 이름 미정'), a) for a in (lst or []))


def etc_ko(o):
    """그 밖의 조건 — 한 줄에 하나씩."""
    out = []
    if 'mode' in o:       out.append('지급 단위 : ' + MODE.get(o['mode'], o['mode']))
    if 'grade' in o:      out.append('수술 종 판정 : ' + GRADE.get(o['grade'], o['grade']))
    if 'cause' in o:      out.append('원인 한정 : %s만' % o['cause'])
    if o.get('ex') == ['cancer']: out.append('암·유사암은 대상에서 제외')
    if 'need_cnt' in o:   out.append('치료를 %d회 이상 받은 단계에서만' % o['need_cnt'])
    if 'need_drug' in o:  out.append('약물 종류가 %d종 이상일 때만' % o['need_drug'])
    if o.get('need_surg'): out.append('같은 단계에 수술이 있어야 지급')
    if o.get('plus'):     out.append('(plus) 담보 — 수술 2회 이상일 때만')
    if o.get('stage') == 'recur': out.append('재진단 단계에서만 (첫 진단 단계에서는 계산 안 함)')
    if 'surg7_grp' in o:  out.append('1-7종 분류표의 수술구분이 「%s」일 때만' % o['surg7_grp'])
    if 'grp' in o:        out.append('질병코드 그룹표를 「%s」로 지정' % o['grp'])
    if o.get('by_benefit'): out.append('세부급부 이름에서 치료행위·암종을 다시 읽음')
    if o.get('hc_listed'): out.append('약관이 진료행위(수가)코드를 열거한 항목은 그 목록에 있을 때만')
    if 'type' in o:       out.append('종류 : ' + o['type'])
    if 'unit' in o:       out.append('지급 방식 : ' + o['unit'])
    if 'log' in o:        out.append('기록 사유 : ' + o['log'])
    if 'note' in o:       out.append('메모 : ' + o['note'])
    return '\n'.join(out)


# 규칙에 적지 않고 담보명 글자에서 자동으로 읽는 조건들 (scen_engine.py 의 tokens·hosp_ok)
TOKEN_DOC = [
    ('상해 · 재해', '원인을 상해로 본다 — 질병 사례에서는 계산하지 않는다', '모든 보상 유형'),
    ('질병', '원인을 질병으로 본다 — 상해 사례에서는 계산하지 않는다', '모든 보상 유형'),
    ('상급종합병원', '상급종합병원에서 받은 치료일 때만. 상급종합은 종합병원 조건도 함께 충족한다',
     '수술비 · 입원일당'),
    ('종합병원', '종합병원 이상에서 받은 치료일 때만', '수술비 · 입원일당'),
    ('요양병원', '요양병원 입원일 때만', '입원일당'),
    ('요양병원제외', '요양병원 입원은 빼고 계산한다 (요양병원을 요구하는 조건이 아니다)', '입원일당'),
    ('1인실', '1인실 입원일 때만 ( "2-3인실" 이 같이 있으면 적용하지 않는다 )', '입원일당'),
    ('2-3인실', '2-3인실 입원일 때만', '입원일당'),
    ('N일한도', '최대 N일까지만 계산한다 (예 : 180일한도)', '입원일당'),
    ('(N일이상', 'N일 이상 입원했을 때부터 계산하고, N일째부터 센다 (예 : (31일이상)', '입원일당'),
    ('[상해2종] · (질병1종)', '수술 종(1~7종)을 담보명에서 읽는다', '수술비'),
    ('중환자실', '중환자실 입원 일수로 계산한다', '입원일당'),
    ('통원', '통원 횟수로 계산한다', '통원일당'),
    ('(plus)', '수술을 2회 이상 받은 단계에서만 계산한다', '수술비'),
    ('특정N대질병제외', '해당 질병군은 대상에서 뺀다', '수술비'),
    ('암종 · 그룹 이름', '담보명에 들어 있는 그룹 이름 중 가장 긴 것으로 질병코드 그룹표를 고른다'
     ' (예 : "10대특정암(전이포함)" 이 "10대특정암" 보다 먼저)', '질병코드 판정'),
    ('[세부급부]', '담보명 끝의 대괄호 안이 보상 유형을 정한다 — 부모 담보명만 보고 단정하지 않는다',
     '규칙 고르기'),
    ('고지유형 꼬리표', '「제외 상품·꼬리표」 시트의 꼬리표는 규칙과 맞춰 보기 전에 떼어 낸다', '규칙 고르기'),
]

def zero_why(rule, cnt, db, S):
    """적용 담보가 0건인 줄이 왜 0건인지 — 이름이 아예 없는 것과,
    위쪽 줄이 먼저 잡아가는 것은 뜻이 전혀 다르다."""
    import re, collections
    if cnt.get(rule['id'], 0):
        return ''
    rx = re.compile(rule['m'])
    xr = re.compile(rule['x']) if rule.get('x') else None
    taken = collections.Counter()
    for r in db['riders']:
        nm = S.nname(r['n'])
        if not rx.search(nm) or (xr and xr.search(nm)):
            continue
        h = S.classify(r['n'])
        taken[h['id'] if h else '(규칙 없음)'] += 1
    if not taken:
        return '특약 마스터에 이 담보명이 없음 — 상품에 없는 담보이거나 이름이 달라진 것'
    top = taken.most_common(2)
    s = '이름이 걸리는 담보 %d건이 있으나 위쪽 줄이 먼저 잡아감 : ' % sum(taken.values())
    s += ' · '.join('「%s」 %d건' % (k, v) for k, v in top)
    return s + '\n→ 이 줄이 실제로 쓰이는지, 순서를 올려야 하는지 확인 필요'


COLS = [
    ('순번', 6, '위에서부터 검사하는 순서. 먼저 걸린 규칙 하나만 적용된다'),
    ('규칙 ID', 20, '코드에서 쓰는 이름'),
    ('보상 유형', 19, '이 담보를 어떤 보상으로 계산할지'),
    ('담보명 조건', 46, '담보명이 이 조건에 걸리면 이 규칙을 쓴다 (| 는 "또는")'),
    ('제외 조건', 26, '담보명이 여기에 걸리면 이 규칙을 쓰지 않고 다음 줄로 넘어간다'),
    ('질병코드(KCD) 판정', 34, '고객의 질병코드를 무엇과 대조해 지급 여부를 정하는가'),
    ('사례 금액 계산', 13, '치료사례 지면에서 금액을 계산하는 유형인가 (아니면 담보 분류만)'),
    ('지급 횟수', 12, '한 사례에서 몇 번까지 지급하는가'),
    ('필요한 치료행위', 24, '사례에서 이 치료를 받았어야 지급 (하나라도 해당되면)'),
    ('모두 필요한 치료행위', 18, '나열된 치료를 모두 받았을 때만 지급'),
    ('질병 계열', 11, '사례가 선언한 질병 계열 한정'),
    ('그 밖의 조건', 34, '위 칸으로 안 적히는 조건들'),
    ('지면 묶음', 12, '제안서에서 어느 묶음으로 보여주는가'),
    ('지급 사유 문구', 34, '제안서에 "왜 지급되는가"로 인쇄되는 문장'),
    ('규칙 설명(메모)', 40, '규칙을 넣은 이유·근거 (코드 주석)'),
    ('적용 담보 수', 11, '특약 마스터 1,757건 중 이 규칙에 걸리는 담보 수'),
    ('0건인 이유', 40, '적용 담보가 0건인 줄만 채워집니다 — 이름이 없는 것인지, 위쪽 줄이 먼저 잡아가는 것인지'),
    ('적용 담보 예시', 40, '실제로 걸린 담보 이름 (최대 5건)'),
    ('검수 결과', 11, '검수자가 적는 칸 — O / X / 보류'),
    ('검수 의견', 34, '검수자가 적는 칸'),
]


def style_head(ws, row=1, ncol=None):
    ncol = ncol or ws.max_column
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, size=9.5, bold=True, color='FFFFFF')
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 30


def build():
    doc = json.load(open(os.path.join(BASE, 'rules.json'), encoding='utf-8'))
    rules = doc['rules']

    # 특약 마스터를 실제로 규칙표에 통과시켜 규칙별 적용 담보를 센다
    sys.path.insert(0, BASE)
    import scen_engine as S
    db = json.load(open(os.path.join(BASE, 'db.json'), encoding='utf-8'))
    cnt = collections.Counter()
    samp = collections.defaultdict(list)
    for r in db['riders']:
        hit = S.classify(r['n'])
        key = hit['id'] if hit else '(규칙 없음)'
        cnt[key] += 1
        if len(samp[key]) < 5:
            samp[key].append(r['n'])
    total_riders = len(db['riders'])

    wb = Workbook()

    # ═══ 읽는 법 ═══
    ws = wb.active
    ws.title = '읽는 법'
    ws.sheet_view.showGridLines = False
    LINES = [
        ('t', '보험금 지급 규칙표 — 검수용'),
        ('s', '메리츠 스마트 제안서 생성기 v8.49 · rules.json 의 규칙 %d줄' % len(rules)),
        ('', ''),
        ('h', '이 표가 하는 일'),
        ('p', '고객 설계서에서 담보명을 하나 읽으면, 이 표를 위에서부터 훑어 처음 걸리는 한 줄로 그 담보를 계산합니다.'),
        ('p', '두 줄에 다 걸려도 위에 있는 줄만 쓰기 때문에 「순번」이 곧 우선순위입니다. 순서를 바꾸면 계산이 바뀝니다.'),
        ('p', '어느 줄에도 걸리지 않으면 금액을 만들지 않고 사유만 남깁니다 — 틀린 금액이 아니라 빈칸으로 나타납니다.'),
        ('', ''),
        ('h', '무엇을 봐 주셔야 하나'),
        ('p', '1. 「담보명 조건」이 우리 담보를 제대로 집어내는지 — 빠진 담보, 엉뚱하게 걸리는 담보'),
        ('p', '2. 「지급 횟수」가 약관과 맞는지 — 최초 1회인데 연 1회로 적혀 있지 않은지'),
        ('p', '3. 「그 밖의 조건」이 약관 조건을 빠뜨리지 않았는지 — 병원 종별, 며칠 이상, 몇 종 이상'),
        ('p', '4. 「적용 담보 수」가 0인 줄 — 지금 마스터에 걸리는 담보가 없는 규칙입니다'),
        ('p', '5. 「지급 사유 문구」가 고객에게 그대로 인쇄되는 문장이므로 표현이 맞는지'),
        ('p', '※ 병원 종별·1인실·N일 이상·N일 한도·수술 종 같은 조건은 규칙에 적지 않습니다 — 담보명 글자에서 자동으로 읽습니다.'),
        ('p', '   그래서 「그 밖의 조건」이 비어 있어도 조건이 없는 것이 아닙니다. 「담보명에서 읽는 조건」 시트를 함께 봐 주세요.'),
        ('', ''),
        ('h', '적어 주시는 칸'),
        ('p', '맨 오른쪽 두 칸(「검수 결과」·「검수 의견」)만 비어 있습니다. 나머지는 프로그램이 읽어 쓴 값이라 여기서 고쳐도 계산에 반영되지 않습니다.'),
        ('p', '고쳐야 할 것은 의견 칸에 적어 주시면 규칙 파일에 반영하겠습니다.'),
        ('', ''),
        ('h', '질병코드(KCD)에 대하여'),
        ('p', '질병코드는 어떤 경우에도 프로그램이 만들어내지 않습니다. 약관 별표에서 뽑아 둔 목록(특약 마스터 · 그룹표)만 씁니다.'),
        ('p', '그래서 「질병코드 판정」이 "…와 대조"로 적힌 줄은, 그 대조표에 코드가 없으면 지급하지 않고 사유를 남깁니다.'),
        ('', ''),
        ('h', '시트 안내'),
        ('p', '규칙 %d개 — 규칙표 본문. 한 줄이 규칙 하나입니다.' % len(rules)),
        ('p', '요약 — 보상 유형별·지급 횟수별 줄 수 (본문을 세는 수식으로 되어 있어 본문을 고치면 함께 바뜁니다).'),
        ('p', '치료행위 어휘 — 「필요한 치료행위」 칸에 쓰는 약어와 한글 이름.'),
        ('p', '질병코드 그룹표 — 「…그룹표와 대조」로 판정하는 줄이 참조하는 분류표 %d종.' % (len([g for g in doc['kcd_groups'] if not g.startswith('_')])),),
        ('p', '담보명에서 읽는 조건 — 규칙에 안 적혀 있고 담보명 글자에서 자동으로 읽는 조건들.'),
        ('p', '제외 상품·고지 꼬리표 — 지면을 붙이지 않는 상품, 담보명 뒤에서 떼어 내는 꼬리표.'),
    ]
    r = 2
    for kind, s in LINES:
        c = ws.cell(row=r, column=2, value=s)
        if kind == 't':   c.font = Font(name=FONT, size=18, bold=True, color='161A1F')
        elif kind == 's': c.font = Font(name=FONT, size=10, color='6B7280')
        elif kind == 'h': c.font = Font(name=FONT, size=11.5, bold=True, color=RED)
        else:             c.font = Font(name=FONT, size=10, color='3A3F45')
        c.alignment = Alignment(vertical='center')
        ws.row_dimensions[r].height = 30 if kind == 't' else (21 if kind == 'h' else 17)
        r += 1
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 132

    # ═══ 규칙 본문 ═══
    name = '규칙 %d개' % len(rules)
    ws = wb.create_sheet(name)
    for i, (h, w, note) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.comment = None
        ws.column_dimensions[get_column_letter(i)].width = w
        ws.cell(row=2, column=i, value=note)
    style_head(ws, 1)
    for c in range(1, len(COLS) + 1):
        cell = ws.cell(row=2, column=c)
        cell.font = Font(name=FONT, size=8, color='8A9099', italic=True)
        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        cell.fill = PatternFill('solid', fgColor='F4F6F8')
        cell.border = BOX
    ws.row_dimensions[2].height = 34

    for i, x in enumerate(rules, 1):
        o = x.get('opt') or {}
        row = [
            i,
            x['id'],
            KIND.get(x['kind'], x['kind']),
            x['m'],
            x.get('x', ''),
            KCD.get(o.get('kcd'), str(o['kcd'])) if 'kcd' in o else KCD['none'],
            '계산함' if x['kind'] in CALC else '분류만 (금액 계산 안 함)',
            FREQ.get(o['freq']) if 'freq' in o else
            (FREQ[FREQ_DEFAULT[x['kind']]] + ' (기본값)' if x['kind'] in FREQ_DEFAULT else '—'),
            acts_ko(o.get('acts')),
            acts_ko(o.get('acts_all')),
            ' · '.join(FAM.get(f, f) for f in (o.get('fam') or [])),
            etc_ko(o),
            o.get('group', ''),
            o.get('why', ''),
            (x.get('label') or '') + (('\n[' + x['_v83'] + ']') if x.get('_v83') else ''),
            cnt.get(x['id'], 0),
            zero_why(x, cnt, db, S),
            '\n'.join(samp.get(x['id'], [])),
            '', '',
        ]
        rr = i + 2
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=rr, column=c, value=v)
            cell.font = Font(name=FONT, size=9,
                             color='161A1F' if c in (2, 3, 4) else '3A3F45',
                             bold=(c == 3))
            cell.alignment = Alignment(
                horizontal='center' if c in (1, 7, 8, 16) else 'left',
                vertical='top', wrap_text=True)
            cell.border = BOX
            if i % 2 == 0:
                cell.fill = BAND
            if c in (19, 20):
                cell.fill = FILLIN
            if c in (16, 17) and cnt.get(x['id'], 0) == 0:
                cell.fill = WARN
                cell.font = Font(name=FONT, size=9, bold=True, color=RED)
        ws.row_dimensions[rr].height = None

    dv = DataValidation(type='list', formula1='"O,X,보류"', allow_blank=True)
    dv.prompt = 'O = 약관과 맞음 / X = 고쳐야 함 / 보류 = 확인 필요'
    dv.promptTitle = '검수 결과'
    ws.add_data_validation(dv)
    dv.add('S3:S%d' % (len(rules) + 2))

    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = 'A2:T%d' % (len(rules) + 2)

    # ═══ 요약 (본문을 세는 수식) ═══
    ws2 = wb.create_sheet('요약')
    ws2.sheet_view.showGridLines = False
    q = "'%s'" % name
    ws2['B2'] = '규칙표 요약'
    ws2['B2'].font = Font(name=FONT, size=15, bold=True)
    ws2['B3'] = ('아래 숫자는 본문 시트를 세는 수식입니다 — 본문을 고치면 함께 바뀝니다. '
                 '칸이 비어 보이면 엑셀에서 파일을 열면 자동으로 계산됩니다.')
    ws2['B3'].font = Font(name=FONT, size=9, color='6B7280')
    last = len(rules) + 2

    def block(r0, title, header, items, col, extra=None):
        ws2.cell(row=r0, column=2, value=title).font = Font(name=FONT, size=11.5, bold=True, color=RED)
        for j, h in enumerate(header):
            c = ws2.cell(row=r0 + 1, column=2 + j, value=h)
            c.font = Font(name=FONT, size=9.5, bold=True, color='FFFFFF')
            c.fill = HEAD
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BOX
        for k, lab in enumerate(items):
            rr = r0 + 2 + k
            a = ws2.cell(row=rr, column=2, value=lab)
            b = ws2.cell(row=rr, column=3,
                         value='=COUNTIF(%s!%s3:%s%d,B%d&"*")' % (q, col, col, last, rr))
            cells = [a, b]
            if extra:
                cells.append(ws2.cell(row=rr, column=4,
                                      value='=SUMIF(%s!%s3:%s%d,B%d&"*",%s!$P$3:$P$%d)' % (q, col, col, last, rr, q, last)))
            for c in cells:
                c.font = Font(name=FONT, size=9.5)
                c.alignment = Alignment(horizontal='left' if c.column == 2 else 'center')
                c.border = BOX
        rr = r0 + 2 + len(items)
        a = ws2.cell(row=rr, column=2, value='합계')
        b = ws2.cell(row=rr, column=3, value='=SUM(C%d:C%d)' % (r0 + 2, rr - 1))
        cells = [a, b]
        if extra:
            cells.append(ws2.cell(row=rr, column=4, value='=SUM(D%d:D%d)' % (r0 + 2, rr - 1)))
        for c in cells:
            c.font = Font(name=FONT, size=9.5, bold=True)
            c.fill = WARN
            c.alignment = Alignment(horizontal='left' if c.column == 2 else 'center')
            c.border = BOX
        return rr + 2

    y = block(5, '보상 유형별', ['보상 유형', '규칙 줄 수', '적용 담보 수'],
              [KIND[k] for k in ('dx', 'surg', 'day', 'tx', 'care', 'life', 'ls', 'nonmed', 'skip', 'itc_unknown')],
              'C', extra=True)
    y = block(y, '지급 횟수별', ['지급 횟수', '규칙 줄 수', '적용 담보 수'],
              [FREQ[k] for k in ('once', 'year', 'each')] + ['—'], 'H', extra=True)
    y = block(y, '질병코드 판정 방식별', ['질병코드(KCD) 판정', '규칙 줄 수', '적용 담보 수'],
              [KCD[k] for k in ('major', 'sim', 'major_or_sim', 'codes', 'g131', 'group', 'hc', 'none')],
              'F', extra=True)

    ws2.cell(row=y, column=2, value='마스터 대조').font = Font(name=FONT, size=11.5, bold=True, color=RED)
    facts = [('특약 마스터 담보 수 (db.json %s)' % db['meta'].get('version', ''), total_riders),
             ('규칙에 걸린 담보 수', total_riders - cnt.get('(규칙 없음)', 0)),
             ('어느 규칙에도 안 걸린 담보 수', cnt.get('(규칙 없음)', 0)),
             ('적용 담보가 0건인 규칙 줄 수', sum(1 for x in rules if cnt.get(x['id'], 0) == 0))]
    for k, (lab, v) in enumerate(facts):
        rr = y + 1 + k
        a = ws2.cell(row=rr, column=2, value=lab)
        b = ws2.cell(row=rr, column=3, value=v)
        for c in (a, b):
            c.font = Font(name=FONT, size=9.5)
            c.alignment = Alignment(horizontal='left' if c.column == 2 else 'center')
            c.border = BOX
    nohit = [r['n'] for r in db['riders'] if S.classify(r['n']) is None]
    ws2.cell(row=y + 1 + len(facts) + 1, column=2,
             value='안 걸린 담보 : ' + (' / '.join(nohit) if nohit else '없음')).font = \
        Font(name=FONT, size=9, color='8E1319')
    ws2.column_dimensions['B'].width = 52
    ws2.column_dimensions['C'].width = 14
    ws2.column_dimensions['D'].width = 14

    # ═══ 치료행위 어휘 ═══
    ws3 = wb.create_sheet('치료행위 어휘')
    used = set()
    for x in rules:
        o = x.get('opt') or {}
        used |= set(o.get('acts') or []) | set(o.get('acts_all') or [])
    ws3.append(['약어', '한글 이름', '규칙에서 쓰는가'])
    style_head(ws3, 1)
    for a in sorted(ACTS):
        ws3.append([a, ACTS[a], '씀' if a in used else '아직 안 씀'])
    for rr in range(2, ws3.max_row + 1):
        for c in range(1, 4):
            cell = ws3.cell(row=rr, column=c)
            cell.font = Font(name=FONT, size=9.5,
                             color='3A3F45' if ws3.cell(row=rr, column=3).value == '씀' else 'A6ACB2')
            cell.border = BOX
            cell.alignment = Alignment(horizontal='center' if c == 3 else 'left')
    for col, w in zip('ABC', (16, 34, 16)):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = 'A2'

    # ═══ 질병코드 그룹표 ═══
    ws4 = wb.create_sheet('질병코드 그룹표')
    ws4.append(['그룹 이름', '코드 수', '질병코드 (약관 별표에서 뽑은 것)'])
    style_head(ws4, 1)
    for g, lst in doc['kcd_groups'].items():
        if g.startswith('_'):
            continue
        ws4.append([g, len(lst), ', '.join(lst)])
    for rr in range(2, ws4.max_row + 1):
        for c in range(1, 4):
            cell = ws4.cell(row=rr, column=c)
            cell.font = Font(name=FONT, size=9.5, color='3A3F45', bold=(c == 1))
            cell.border = BOX
            cell.alignment = Alignment(horizontal='center' if c == 2 else 'left',
                                       vertical='top', wrap_text=(c == 3))
    for col, w in zip('ABC', (30, 10, 110)):
        ws4.column_dimensions[col].width = w
    ws4.freeze_panes = 'A2'

    # ═══ 담보명에서 자동으로 읽는 조건 ═══
    ws6 = wb.create_sheet('담보명에서 읽는 조건')
    ws6.append(['담보명에 이 글자가 있으면', '이렇게 읽는다', '어디에 쓰이는가'])
    style_head(ws6, 1)
    for row in TOKEN_DOC:
        ws6.append(list(row))
    for rr in range(2, ws6.max_row + 1):
        for c in range(1, 4):
            cell = ws6.cell(row=rr, column=c)
            cell.font = Font(name=FONT, size=9.5, color='3A3F45', bold=(c == 1))
            cell.border = BOX
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    for col, w in zip('ABC', (30, 58, 30)):
        ws6.column_dimensions[col].width = w
    ws6.freeze_panes = 'A2'

    # ═══ 제외 상품 · 고지 꼬리표 ═══
    ws5 = wb.create_sheet('제외 상품·꼬리표')
    ws5.sheet_view.showGridLines = False
    ws5['B2'] = '지면을 붙이지 않는 상품 (정책으로 정한 예외)'
    ws5['B2'].font = Font(name=FONT, size=11.5, bold=True, color=RED)
    ws5['B3'] = '상품명이 조건에 걸리면 계산하지 않고 원본 설계서를 그대로 돌려줍니다.'
    ws5['B3'].font = Font(name=FONT, size=9, color='6B7280')
    for j, h in enumerate(['상품 이름', '상품명 조건', '왜 제외하는가']):
        c = ws5.cell(row=5, column=2 + j, value=h)
        c.font = Font(name=FONT, size=9.5, bold=True, color='FFFFFF')
        c.fill = HEAD
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = BOX
    for k, e in enumerate(doc['exclude_products']):
        rr = 6 + k
        for j, v in enumerate([e.get('nm', ''), e.get('m', ''), e.get('why', '')]):
            c = ws5.cell(row=rr, column=2 + j, value=v)
            c.font = Font(name=FONT, size=9.5, color='3A3F45')
            c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            c.border = BOX
    y = 6 + len(doc['exclude_products']) + 2
    ws5.cell(row=y, column=2, value='담보명 뒤에서 떼어 내는 고지유형 꼬리표').font = \
        Font(name=FONT, size=11.5, bold=True, color=RED)
    ws5.cell(row=y + 1, column=2,
             value='설계서 담보명 뒤에 붙는 꼬리표입니다. 담보명을 규칙표와 맞춰 보기 전에 떼어 냅니다.').font = \
        Font(name=FONT, size=9, color='6B7280')
    for k, t in enumerate(doc['goji_tags']):
        c = ws5.cell(row=y + 3 + k, column=2, value=t)
        c.font = Font(name=FONT, size=9.5, color='3A3F45')
        c.border = BOX
    ws5.column_dimensions['B'].width = 34
    ws5.column_dimensions['C'].width = 34
    ws5.column_dimensions['D'].width = 90

    # 같은 파일에 다른 세션이 덧붙여 둔 시트(예 「2차 검수 · 검사비 항목」)는 살려 둔다.
    # 이 스크립트가 만드는 시트만 갈아 끼우고 나머지는 그대로 옮긴다.
    mine = set(wb.sheetnames)
    if os.path.exists(OUT):
        try:
            from openpyxl import load_workbook
            old = load_workbook(OUT)
            for nm in old.sheetnames:
                if nm in mine: continue
                src = old[nm]; dst = wb.create_sheet(nm)
                for row in src.iter_rows():
                    for c in row:
                        if c.value is not None: dst.cell(row=c.row, column=c.column, value=c.value)
                for k, d in src.column_dimensions.items(): dst.column_dimensions[k].width = d.width
                print('  기존 시트 유지 :', nm)
        except Exception as ex:
            print('  ! 기존 파일의 시트를 옮기지 못했습니다(%s) — 새로 만듭니다' % type(ex).__name__)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print('저장 %s · 규칙 %d줄 · 담보 대조 %d건' % (OUT, len(rules), total_riders))


if __name__ == '__main__':
    build()
