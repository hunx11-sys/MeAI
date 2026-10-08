# 이전 설계도 엑셀의 손으로 적은 「⑫ 검수 메모」 시트를 새 엑셀에 옮긴다(build_ga_xlsx 는 엑셀을 새로 만들어 이 시트가 비어 버린다)
import sys
from openpyxl import load_workbook
from copy import copy
old, new = sys.argv[1], sys.argv[2]
wo = load_workbook(old); wn = load_workbook(new)
T = '⑫ 검수 메모'
so = wo[T]
if T in wn.sheetnames:
    idx = wn.sheetnames.index(T); wn.remove(wn[T])
else:
    idx = len(wn.sheetnames)
sn = wn.create_sheet(T, index=idx)
for row in so.iter_rows():
    for c in row:
        d = sn.cell(row=c.row, column=c.column, value=c.value)
        if c.has_style:
            d.font = copy(c.font); d.fill = copy(c.fill); d.border = copy(c.border); d.alignment = copy(c.alignment); d.number_format = c.number_format
for k, dim in so.column_dimensions.items(): sn.column_dimensions[k].width = dim.width
for k, dim in so.row_dimensions.items(): sn.row_dimensions[k].height = dim.height
sn.freeze_panes = so.freeze_panes
wn.save(new); print('restored', T, so.max_row, 'rows ->', new)
