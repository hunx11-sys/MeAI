"""GA 칸별 매핑 엑셀에 '쪽별 캡처 + 칸별 특약 이름' 시트를 붙인다.
  python3 build_ga_visual.py <ga_spec_results.json> <설계서 GA html> <설계서 pdf 이름> <xlsx>
- html 의 금액 칸을 ga_spec 의 칸 순서(토큰 순서)대로 짚어 칸 번호 배지를 달고 렌더 → 쪽별 PNG
- 시트마다 왼쪽 캡처, 오른쪽 칸 번호별 : 이 설계서에서 계산된 특약(금액) · 들어올 수 있는 특약 전체(마스터)
- 기존 시트는 그대로 두고 '캡처 ' 로 시작하는 시트만 다시 만든다."""
import os, sys, os, re, json, collections, asyncio, pathlib
import pymupdf
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XImg
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from playwright.async_api import async_playwright

RES = json.load(open(sys.argv[1], encoding='utf-8'))
HTML, DESIGN, XLSX = sys.argv[2], sys.argv[3], sys.argv[4]
WORK = os.path.join(os.path.dirname(os.path.abspath(HTML)), '_visual'); os.makedirs(WORK, exist_ok=True)
D = next(d for d in RES['designs'] if os.path.basename(d['pdf']) == DESIGN)
CELLS = RES['cells']
PAGES = {1: '보장 한눈에', 2: '암 세부', 3: '뇌심 세부', 4: '암치료 세부내역', 5: '뇌·심장 치료 세부내역', 6: '주요 수술비 세부내역'}

# ── 1. 토큰 → 칸 번호 ──
order = []
for c in CELLS:
    for _ in D['cells'][c['id']]['tokens']: order.append(c['id'])
h = open(HTML, encoding='utf-8').read()
PAT = re.compile(r'<b class="v[^"]*">([^<]+)<i>만원</i></b>|<span class="na">(미가입|면책)</span>|<span class="txt">(지원가능)</span>')
ms = list(PAT.finditer(h))
assert len(ms) == len(order), (len(ms), len(order))
out, last, seen = [], 0, set()
for m, cid in zip(ms, order):
    out.append(h[last:m.start()]); tag = m.group(0)
    if cid not in seen:
        seen.add(cid); tag = re.sub(r'^<(b|span) ', r'<\1 data-cid="%d" ' % int(cid.split('-')[1]), tag)
    out.append(tag); last = m.end()
out.append(h[last:])
# 배지는 레이아웃이 끝난 뒤 쪽(.page) 위에 절대 위치로 얹는다(칸 안에 넣으면 잘림)
JS = """<script>window.__badge=()=>{document.querySelectorAll('[data-cid]').forEach(e=>{const p=e.closest('.page');const P=p.getBoundingClientRect();
 const r=e.getBoundingClientRect();const s=document.createElement('span');s.textContent=e.dataset.cid;
 s.style.cssText='position:absolute;background:#1F4E9E;color:#fff;font:700 7pt/1.3 sans-serif;border-radius:6px;padding:0 3px;z-index:99;box-shadow:0 0 0 1.2px #fff';
 p.appendChild(s);const w=s.getBoundingClientRect().width;s.style.left=(r.left-P.left-w-2)+'px';s.style.top=(r.top-P.top+(r.height-11)/2)+'px';});}</script>"""
out.append(JS)
ah = os.path.join(WORK, 'annotated.html'); open(ah, 'w', encoding='utf-8').write(''.join(out))
apdf = os.path.join(WORK, 'annotated.pdf')

async def render():
    async with async_playwright() as pw:
        b = await pw.chromium.launch(); pg = await b.new_page(viewport={'width': 794, 'height': 1123})
        await pg.goto(pathlib.Path(ah).resolve().as_uri()); await pg.evaluate('document.fonts.ready'); await pg.wait_for_timeout(500); await pg.evaluate('window.__badge()')
        await pg.pdf(path=apdf, prefer_css_page_size=True, print_background=True); await b.close()
asyncio.run(render())
doc = pymupdf.open(apdf); PNG = {}; CALC = []
kinds = re.findall(r'<div class="page( calcpg)?">', h)          # 지면 순서 : 1쪽 → 합산 계산서(calcpg) → 2~6쪽 (v8.68~)
assert len(kinds) == doc.page_count, (len(kinds), doc.page_count)
pno = 0
for i, k in enumerate(kinds):
    p = os.path.join(WORK, 'page%d.png' % (i + 1)); doc[i].get_pixmap(dpi=110).save(p)
    if k: CALC.append(p)
    else: pno += 1; PNG[pno] = p
    if not k:                                                      # 칸별 특약 지도(HTML)용 캡처
        doc[i].get_pixmap(dpi=96).save(os.path.join(WORK, 'p%d.jpg' % pno), jpg_quality=72)

# ── 2. 칸별 특약 ──
def used(lines):
    return [l for l in lines if l.get('counted') and l.get('in_row') is not False and l.get('amt')]
def won(v):
    v = int(round(v)); e, m = divmod(v, 10000)
    return ('%d억 %s' % (e, format(m, ',')) if m else '%d억' % e) if e else format(v, ',')
def design_txt(cid):
    ls = used(D['cells'][cid]['lines'])
    if not ls: return ''
    return '\n'.join('· %s  %s만원' % (l['name'], won(l['amt'])) for l in ls)
def master_txt(cid):
    u = RES['universe'].get(cid)
    if not u: return ''
    prod = collections.OrderedDict()
    for l in used(u['lines']): prod.setdefault(l['name'], set()).add(l.get('product') or '')
    return '\n'.join('· %s%s' % (n, ('  〔%s〕' % '·'.join(sorted(p for p in ps if p))) if any(ps) else '') for n, ps in prod.items())
def master_n(cid):
    u = RES['universe'].get(cid)
    return len({l['name'] for l in used(u['lines'])}) if u else 0
def where(c):
    parts = [c['section']]
    if c['row'] and c['row'] != c['section']: parts.append(c['row'])
    if c['col']: parts.append(c['col'])
    return ' › '.join(p for p in parts if p)
def note(c):
    if c['kind'] == 'engine': return c.get('sc_name') or ''
    return c.get('note') or ('설계서 담보명·가입금액을 그대로 옮기는 칸' if c['kind'] == 'direct' else '')

# ── 3. 시트 ──
wb = load_workbook(XLSX)
for ws in [w for w in wb.worksheets if w.title.startswith('캡처 ')]: wb.remove(ws)
HF = PatternFill('solid', fgColor='1F4E9E'); ZF = PatternFill('solid', fgColor='F3F6FB')
thin = Side(style='thin', color='C9CED6'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
meta = D.get('meta') or {}
for pno in range(1, 7):
    ws = wb.create_sheet('캡처 %d쪽 %s' % (pno, PAGES[pno])[:31], index=pno)
    ws.sheet_view.showGridLines = False
    ws['A1'] = 'GA 스마트 제안서 %d쪽 「%s」 — 칸 번호별로 계산되는 특약' % (pno, PAGES[pno]); ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = ('캡처는 설계서 %s 로 만든 지면(파란 숫자 = 칸 번호). 오른쪽 표의 「이 설계서에서 계산된 특약」은 그 설계서 기준 실제 금액, '
                '「들어올 수 있는 특약 전체」는 특약 마스터 전체를 같은 사례에 태웠을 때 이 칸에 잡히는 특약 이름(〔 〕= 상품). '
                '특약별 산정 기준은 ④ 시트, 계산 제외 사유는 ⑧ 시트.' % DESIGN)
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells('A2:P2'); ws.row_dimensions[2].height = 32
    for col in 'ABCDEFGH': ws.column_dimensions[col].width = 11
    ws.column_dimensions['I'].width = 2
    img = XImg(PNG[pno]); ratio = img.height / img.width; img.width = 640; img.height = int(640 * ratio)
    ws.add_image(img, 'A5')
    hdr = ['칸 번호', '위치(구역 › 줄 › 칸)', '지면 값(만원)', '이 설계서에서 계산된 특약 · 금액', '들어올 수 있는 특약 전체(마스터) 수', '들어올 수 있는 특약 전체(마스터)', '사례(가정) · 비고']
    widths = [8, 30, 12, 58, 11, 62, 36]
    for j, (t, w) in enumerate(zip(hdr, widths)):
        cell = ws.cell(row=4, column=10 + j, value=t); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = HF
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); cell.border = BD
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[4].height = 30
    r = 5
    for k, c in enumerate([c for c in CELLS if c['page'] == pno]):
        dv = design_txt(c['id']); mv = master_txt(c['id'])
        iss = D['cells'][c['id']].get('issues') or []
        if iss:
            dv = (dv + '\n' if dv else '') + '\n'.join('※ [%s] %s — %s' % (x.get('구분', ''), x.get('담보', ''), x.get('사유', '')) for x in iss[:8]) + ('\n※ 외 %d건(⑧ 시트)' % (len(iss) - 8) if len(iss) > 8 else '')
        vals = [int(c['id'].split('-')[1]), where(c), ' / '.join(D['cells'][c['id']]['tokens']),
                dv or ('(계산된 특약 없음 — 미가입·면책·0)' if c['kind'] == 'engine' else '(계산 칸 아님)'), master_n(c['id']) if c['kind'] == 'engine' else '', mv, note(c)]
        for j, v in enumerate(vals):
            cell = ws.cell(row=r, column=10 + j, value=v); cell.alignment = WRAP; cell.border = BD
            if k % 2: cell.fill = ZF
        ws.cell(row=r, column=10).alignment = Alignment(horizontal='center', vertical='top')
        ws.cell(row=r, column=10).font = Font(bold=True, color='1F4E9E')
        ws.cell(row=r, column=12).alignment = Alignment(horizontal='right', vertical='top')
        n = max(dv.count('\n') + 1, min(mv.count('\n') + 1, 40), 1)
        ws.row_dimensions[r].height = min(409, max(18, 14.5 * n))
        r += 1
    ws.freeze_panes = 'A5'
    print(pno, PAGES[pno], 'cells', r - 5)
for j, p in enumerate(CALC, 1):                                   # 1쪽 뒤 합산 계산서 — 그림만(내용은 ⑪ 설계서 칸별 특약 금액 시트와 같다)
    ws = wb.create_sheet('캡처 1-%d 합산 계산서' % j, index=1 + j)
    ws.sheet_view.showGridLines = False
    ws['A1'] = '1쪽 「보장 한눈에」 뒤에 붙는 합산 계산서 %d/%d — 1쪽 칸마다 어떤 설계 특약을 더해 그 금액이 나왔는지(특약명 · 가입금액 · 지급액). 설계서 %s' % (j, len(CALC), DESIGN); ws['A1'].font = Font(bold=True, size=13)
    img = XImg(p); ratio = img.height / img.width; img.width = 900; img.height = int(900 * ratio); ws.add_image(img, 'A3')
ov = wb.worksheets[0]; MARK = '【캡처 시트】'
if not any(isinstance(c.value, str) and c.value.startswith(MARK) for row in ov.iter_rows() for c in row):
    r0 = ov.max_row + 2
    ov.cell(row=r0, column=1, value=MARK + ' 「캡처 1쪽~6쪽」 시트 — 지면 캡처에 칸 번호(파란 숫자)를 달고, 옆 표에 칸 번호별로 계산되는 특약 이름·금액을 나열. 「캡처 1-①~③ 합산 계산서」는 1쪽 뒤에 반드시 붙는 계산서 지면(GA 확정 2026-10-02). '
            '먼저 이 시트로 칸을 찾고, 자세한 조건은 ②·③·④ 시트에서 같은 칸 번호(P쪽-번호)로 찾는다.').font = Font(bold=True, color='1F4E9E')
wb.save(XLSX)
print('saved', XLSX, [w.title for w in wb.worksheets])
